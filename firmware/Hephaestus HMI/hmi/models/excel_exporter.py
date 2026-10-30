import datetime as dt
import math

import numpy as np
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill

from . import registers

HEADERS = (
    ("Date / Time", 21),
    ("Elapsed (s)", 12),
    ("Link", 11),
    ("Run state", 11),
    ("Control mode", 14),
    ("Temperature (°C)", 17),
    ("Setpoint (°C)", 14),
    ("Heater output (%)", 17),
    ("Manual output (%)", 17),
    ("Kp", 10),
    ("Ki", 10),
    ("Kd", 10),
    ("System state", 14),
    ("Fault code", 11),
    ("Fault description", 24),
)

DATA_COLUMNS = 12


def _number(value):
    return None if math.isnan(value) else float(value)


def _header_cell(sheet, text):
    cell = WriteOnlyCell(sheet, value=text)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", start_color="1F3A5F")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    return cell


def export_to_excel(path, snapshot, start, end, progress=None):
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Data")
    for index, (_, width) in enumerate(HEADERS):
        sheet.column_dimensions[chr(ord("A") + index)].width = width
    sheet.freeze_panes = "A2"
    sheet.append([_header_cell(sheet, title) for title, _ in HEADERS])

    seconds = np.floor(snapshot.times).astype(np.int64)
    last_index = {int(second): i for i, second in enumerate(seconds)}
    first_second = int(math.floor(start))
    last_second = int(math.floor(end))
    total = last_second - first_second + 1
    missing = 0

    for n, second in enumerate(range(first_second, last_second + 1)):
        stamp = dt.datetime.fromtimestamp(second)
        elapsed = second - first_second
        i = last_index.get(second)
        if i is None or snapshot.data[i, 0] == 0.0:
            missing += 1
            sheet.append([stamp, elapsed, "NO DATA"] + [None] * (len(HEADERS) - 3))
        else:
            row = snapshot.data[i]
            state = int(row[10])
            fault = int(row[11])
            sheet.append(
                [
                    stamp,
                    elapsed,
                    "ONLINE",
                    "RUNNING" if row[1] else "STOPPED",
                    "MANUAL" if row[2] else "AUTO (PID)",
                    _number(row[3]),
                    _number(row[4]),
                    _number(row[5]),
                    _number(row[6]),
                    _number(row[7]),
                    _number(row[8]),
                    _number(row[9]),
                    registers.describe_state(state),
                    fault,
                    registers.describe_fault(fault),
                ]
            )
        if progress is not None and n % 2000 == 0:
            progress(int(100 * n / total))

    info = workbook.create_sheet("Info")
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 26
    info.append([_header_cell(info, "Item"), _header_cell(info, "Value")])
    info.append(["Export start", dt.datetime.fromtimestamp(first_second)])
    info.append(["Export end", dt.datetime.fromtimestamp(last_second)])
    info.append(["Rows (1 per second)", total])
    info.append(["Seconds with no data", missing])
    info.append(["Created", dt.datetime.now()])

    workbook.save(path)
    if progress is not None:
        progress(100)
    return total, missing
